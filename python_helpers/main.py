import lxml.etree as gfg
import pandas as pd
import numpy as np
import openpyxl
  

def generate_clothing_item(name, model, texture_choices, guid = None):
    root = gfg.Element("clothingItem")

    m_MaleModel = gfg.Element("m_MaleModel")
    m_MaleModel.text = f"{model}_Male"
    root.append(m_MaleModel)

    m_FemaleModel = gfg.Element("m_FemaleModel")
    m_FemaleModel.text = f"{model}_Female"
    root.append(m_FemaleModel)

    m_GUID = gfg.Element("m_GUID")
    if guid:
        m_GUID.text = guid
    else:
        m_GUID.text = "get guid from func"

    root.append(m_GUID)

    m_Static = gfg.Element("m_Static")
    m_Static.text = "false"
    root.append(m_Static)

    m_AllowRandomTint = gfg.Element("m_AllowRandomTint")
    m_AllowRandomTint.text = "false"
    root.append(m_AllowRandomTint)

    # Defined by the amount of textures that we're gonna pass
    for tex in texture_choices:
        textureChoices = gfg.Element("textureChoices")
        textureChoices.text = tex
        root.append(textureChoices)

    tree = gfg.ElementTree(root)

    path = r'python_helpers/outputs/output_clothing/' + name + ".xml"

    with open(path, "wb") as file:
        tree.write(file, encoding='utf-8', xml_declaration=True, pretty_print=True )

def generate_recipe(recipe_name, recipe_items, result_name, time, skill_required, tooltip):
    root_element = f"recipe {recipe_name}\n"
    root_element += "\t{\n"

    for item in recipe_items:
        root_element += f"\t\t{item},\n"

    root_element += f"\n\n\t\tResult: {result_name},\n"
    root_element += f"\t\tTime: {time:.2f},\n"
    root_element += "\t\tNeedToBeLearn: true,\n"
    root_element += "\t\tCanBeDoneFromFloor: false,\n"
    root_element += "\t\tOnGiveXP: NoXP_OnGiveXP,\n"
    root_element += f"\t\tSkillRequired: {skill_required[0]}={skill_required[1]},\n"
    root_element += "\t\tCategory: Surgeon,\n"
    root_element += f"\t\tTooltip: {tooltip},\n"

    root_element += "\t}\n"


    path = r'python_helpers/outputs/output_recipe/script.txt'

    with open(path, "at") as file:
        file.write(root_element)
        file.close()

def generate_item(item_name, weight, item_type, display_category, display_name, icon, tooltip, can_have_holes, clothing_item=None, body_location = None, blood_location = None):
    root_element = f"item {item_name}\n"
    root_element += "\t{\n"

    root_element += f"\t\tWeight = {weight},\n"
    root_element += f"\t\tType = {item_type},\n"
    root_element += f"\t\tDisplayCategory = {display_category},\n"
    root_element += f"\t\tDisplayName = {display_name},\n"
    
    if item_type == "Clothing":
        root_element += f"\t\tClothingItem = {clothing_item},\n"
        root_element += f"\t\tBodyLocation = {body_location},\n"
        root_element += f"\t\tBloodLocation = {blood_location},\n"
    
    root_element += f"\t\tIcon = {icon},\n"
    root_element += f"\t\tTooltip = {tooltip},\n"
    root_element += f"\t\tCanHaveHoles = {can_have_holes.lower()},\n"

    root_element += "\t}\n"


    path = r'python_helpers/outputs/output_item/script.txt'

    with open(path, "at") as file:
        file.write(root_element)
        file.close()


###########################################################################################
def read_table(file_name: str, table_name: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(file_name, read_only= False, data_only = True) # openpyxl does not have table info if read_only is True; data_only means any functions will pull the last saved value instead of the formula
    for sheetname in wb.sheetnames: # pulls as strings
        sheet = wb[sheetname] # get the sheet object instead of string
        if table_name in sheet.tables: # tables are stored within sheets, not within the workbook, although table names are unique in a workbook
            tbl = sheet.tables[table_name] # get table object instead of string
            tbl_range = tbl.ref #something like 'C4:F9'
            break # we've got our table, bail from for-loop
    data = sheet[tbl_range] # returns a tuple that contains rows, where each row is a tuple containing cells
    content = [[cell.value for cell in row] for row in data] # loop through those row/cell tuples
    header = content[0] # first row is column headers
    rest = content[1:] # every row that isn't the first is data
    df = pd.DataFrame(rest, columns = header)
    wb.close()
    return df

###########################################################################################

excel_path = r'python_helpers/modules_prost.xlsx'
df_base = read_table(excel_path, "BaseTable")
df_top = read_table(excel_path, "TopTable")


# CLOTHING GENERATION PASS

limbs = ["Hand", "LowerArm"]
sides = ["Left", "Right"]
prost_bodylocations = ["TOC_ArmRightProsthesis", "TOC_ArmLeftProsthesis"]



for base_row in df_base.iterrows():
    for top_row in df_top.iterrows():
        base_name = base_row[1][0]
        top_name = top_row[1][0]

        for limb in limbs:
            for side in sides:
                current_name = "Prost_" + side + "_" + limb + "_" + base_name + "_" + top_name
                generate_clothing_item(current_name, "test", {"test1", "test2"}, "123")






# ITEM GENERATION PASS - ASSEMBLED
for base_row in df_base.iterrows():
    for top_row in df_top.iterrows():
        base_id = base_row[1]["Base"]
        top_id = top_row[1]["Top"]

        item_id = "Prost_" + base_id + "_" + top_id
        item_type = "Clothing"
        weight = "{0:.2f}".format(float(base_row[1]["Weight"]) + float(top_row[1]["Weight"]))
        display_category = "Prosthesis"
        display_name = "Prosthesis - " + base_row[1]["Display Name"] + " and " + top_row[1]["Display Name"]

        for limb in limbs:
            for side in sides:
                clothing_item_name = "Prost_" + side + "_" + limb + "_" + base_id + "_" + top_id
                bl = prost_bodylocations[0] if side == "Right" else prost_bodylocations[1]

                icon = "metalLeg"
                generate_item(item_id, weight, item_type, display_category, display_name, icon, "TempTooltip", "false", clothing_item_name, bl, "TestBloodLocation")


# ITEM GENERATION PASS - Single item to assemble stuff
def generate_normal_items(df, type):
    for row in df.iterrows():
        item_id = "ProstPart_" + row[1][type]
        item_type = "Normal"
        weight = "{0:.2f}".format(float(row[1]["Weight"]))
        display_category = "Prosthesis"
        display_name = row[1]["Display Name"]
        icon = "ProstTest" + type
        generate_item(item_id, weight, item_type, display_category, display_name, icon, "TempTooltip", "false")

generate_normal_items(df_base, "Base")
generate_normal_items(df_top, "Top")


# RECIPE GENERATION PASS
for base_row in df_base.iterrows():
    for top_row in df_top.iterrows():
        recipe_name = f"Craft prosthesis with {base_name} and {top_name}"

        first_item = "ProstPart_" + base_row[1]["Base"]
        second_item = "ProstPart_" + top_row[1]["Top"]

        recipe_items = [first_item, second_item]
        result = "Prost_" + base_row[1]["Base"] + "_" + top_row[1]["Top"]
        time = 10       # TODO Change this
        skill_required = ["FirstAid","2"]       # TODO Change this
        tooltip = "Tooltip_test"
        generate_recipe(recipe_name, recipe_items, result, time, skill_required, tooltip)
